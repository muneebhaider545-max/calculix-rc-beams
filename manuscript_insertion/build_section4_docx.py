#!/usr/bin/env python3
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import csv

OUT=Path("manuscript_insertion")
OUT.mkdir(exist_ok=True)
fig=Path("model_matrix_outputs/Figure_15_12_case_3D_model_matrix.png")
csvp=Path("model_matrix_outputs/12_case_model_matrix.csv")

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.75); sec.bottom_margin=Inches(0.75); sec.left_margin=Inches(0.8); sec.right_margin=Inches(0.8)

styles=doc.styles
styles["Normal"].font.name="Times New Roman"; styles["Normal"].font.size=Pt(12)
for s in ["Title","Heading 1","Heading 2"]:
    styles[s].font.name="Times New Roman"

p=doc.add_paragraph()
r=p.add_run("Replacement Section 4 for the RC-beam strengthening manuscript")
r.bold=True; r.font.name="Times New Roman"; r.font.size=Pt(14)
p.alignment=WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading("4. Numerical modelling framework and 12-case three-dimensional model matrix", level=1)
p=doc.add_paragraph(
"A common three-dimensional model architecture was prepared for all 12 experimental configurations so that the numerical phase follows the same specimen matrix used in the physical testing. The matrix comprises the control, externally bonded steel-sheet, GFRP and CFRP configurations at nominal concrete strengths of 18, 21 and 25 MPa. The experimentally measured concrete compressive strengths of 19.55, 26.90 and 28.70 MPa are retained for the corresponding numerical cases. The model geometry reproduces the 1219.2 mm × 101.6 mm × 203.2 mm beam, the reported internal reinforcement arrangement, the support locations, and centre-point loading. Strengthening layers are placed on the soffit in the strengthened cases only."
)

doc.add_heading("4.1 Model matrix", level=2)
p=doc.add_paragraph(
"The 12 numerical cases are designated C18-CONT, C18-ST, C18-GFRP, C18-CFRP, C21-CONT, C21-ST, C21-GFRP, C21-CFRP, C25-CONT, C25-ST, C25-GFRP and C25-CFRP. A common mesh architecture is used at the model-generation stage to ensure that differences among cases arise from substrate strength and strengthening system rather than arbitrary changes in discretization. The visual model matrix is shown in Figure 15."
)

rows=list(csv.DictReader(open(csvp,newline="")))
table=doc.add_table(rows=1, cols=7)
table.alignment=WD_TABLE_ALIGNMENT.CENTER
table.style="Table Grid"
hdr=table.rows[0].cells
for i,t in enumerate(["Specimen","System","Measured f′c (MPa)","First crack (kN)","Ultimate load (kN)","Peak deflection (mm)","Strengthening geometry"]):
    hdr[i].text=t
for row in rows:
    cells=table.add_row().cells
    geom="—"
    if row["System"]=="CFRP": geom="4 mm × 101.6 mm; full soffit length"
    elif row["System"]=="Steel": geom="4 mm sheet; full soffit length"
    elif row["System"]=="GFRP": geom="990.6 mm bonded length; thickness not reported"
    vals=[row["Specimen"],row["System"],row["Measured_fc_MPa"],row["First_crack_kN"],row["Ultimate_kN"],row["Peak_deflection_mm"],geom]
    for i,v in enumerate(vals):
        cells[i].text=str(v)
        cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

cap=doc.add_paragraph()
cap.add_run("Table 7. ").bold=True
cap.add_run("Twelve-case numerical model matrix aligned with the experimental programme. GFRP thickness is not specified in the source programme and is therefore not asserted as an experimentally measured model input.")
cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

if fig.exists():
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(fig),width=Inches(7.0))
    cp=doc.add_paragraph()
    cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cp.add_run("Figure 15. ").bold=True
    cp.add_run("Three-dimensional finite-element model matrix for the 12 experimental configurations. The panels document geometry, reinforcement, support/load arrangement and the presence or absence of the external strengthening layer; they are model-architecture visualizations rather than solver-result contours.")

doc.add_heading("4.2 Geometry, reinforcement and strengthening layers", level=2)
doc.add_paragraph(
"Concrete is represented as a three-dimensional solid body. Longitudinal and transverse reinforcement are represented explicitly within the beam volume so that the internal reinforcement layout remains visible and reproducible across all cases. The support centres are positioned 76.2 mm from the beam ends, and the load is applied at midspan. CFRP is represented by a continuous soffit layer using the reported 4 mm thickness, 101.6 mm width and 1219.2 mm length. The bonded steel sheet is represented using the reported 4 mm thickness. The GFRP source programme reports a bonded length of 39 in. (990.6 mm), leaving approximately 4.5 in. (114.3 mm) unstrengthened at each end; because the laminate thickness is not reported, no experimentally derived GFRP thickness is claimed in the model matrix."
)

doc.add_heading("4.3 Constitutive and interface strategy", level=2)
doc.add_paragraph(
"For nonlinear analysis, the three measured concrete strengths should be represented by a damage-plasticity or equivalent nonlinear concrete constitutive law, while internal steel reinforcement should use elastic–plastic behaviour. CFRP and GFRP should be represented as orthotropic external reinforcement when complete laminate properties are available, whereas the steel sheet can use the reported elastic modulus, yield strength and tensile strength. The bonded interface should be represented with a cohesive or equivalent traction–separation formulation rather than an unconditional perfect tie when interface failure is being investigated. Parameters that were not measured in the source experiments—particularly cohesive stiffness, interface strengths, fracture energies and GFRP laminate thickness—must be reported explicitly as calibrated or literature-derived inputs."
)

doc.add_heading("4.4 Validation hierarchy and reporting", level=2)
doc.add_paragraph(
"The numerical validation should proceed in two stages. First, the three control beams should be calibrated against their experimental load–deflection response and peak capacity without using the strengthened specimens to tune the concrete model. Second, the steel-, GFRP- and CFRP-strengthened configurations should be analysed using the same calibrated concrete parameters, changing only the strengthening-system and interface definitions. Validation should compare peak load, peak deflection, initial/secondary stiffness, energy absorption and observed failure mode. Mesh sensitivity should be checked independently before quantitative interpretation of interface stresses or damage localization."
)

doc.add_heading("4.5 Status of the present numerical contribution", level=2)
doc.add_paragraph(
"The 12 three-dimensional model geometries have been generated and organized using a single specimen matrix. At this stage, Figure 15 documents the complete model architecture only. Numerical stress, strain, concrete-damage and interface-damage contours should be inserted only after successful execution of the selected nonlinear solver and extraction of native output files. This distinction is retained to prevent model visualizations from being interpreted as computed results."
)

out=OUT/"Section_4_12_Case_Numerical_Model_Matrix.docx"
doc.save(out)

txt=OUT/"Section_4_replacement_text.txt"
txt.write_text("\n".join(p.text for p in doc.paragraphs),encoding="utf-8")
print(out)
